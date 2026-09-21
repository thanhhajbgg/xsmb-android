import csv,io,sys
from pathlib import Path
from xml.sax.saxutils import escape

HEADERS={'a':'Số thứ nhất','b':'Số thứ hai','either_days':'Ngày có ít nhất 1 số','both_days':'Ngày có cả 2 số','sample_days':'Số kỳ mẫu','either_rate':'Tỷ lệ ít nhất 1 (%)','both_rate':'Tỷ lệ cả 2 (%)','reasons':'Cầu hỗ trợ','both':'Cả hai xuất hiện','day':'Ngày','number':'Số lô','points':'Điểm','cost_rate':'Giá mua/điểm','payout_rate':'Trả/điểm','hits':'Nháy',
 'cost':'Tiền mua (đ)','received':'Tiền nhận (đ)','profit':'Lãi/lỗ (đ)','cumulative':'Lũy kế (đ)','source':'Nguồn',
 'note':'Ghi chú','picks':'Số chọn','score':'AI Score','cluster':'Cụm','cash':'Vốn khả dụng (đ)','roi':'ROI (%)',
 'opening':'Vốn đầu kỳ (đ)','pending_cost':'Tiền chờ (đ)','kind':'Loại nhật ký','created_at':'Thời điểm lưu'}


def assets():return Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parent.parent))/'assets'


def scalar(v):
    if isinstance(v,(dict,list,tuple)):
        if isinstance(v,(list,tuple)):return ', '.join(str(x) for x in v)
        return str(v)
    return v


def safe(v):
    v=scalar(v)
    return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@','\t','\r')) else v


def export_bytes(fmt,title,rows,keys,note=''):
    headers=[HEADERS.get(k,k) for k in keys]
    values=[[scalar(r.get(k)) for k in keys] for r in rows]
    if fmt=='csv':
        buf=io.StringIO(newline='');w=csv.writer(buf);w.writerow(headers)
        w.writerows([[safe(v) for v in row] for row in values])
        return buf.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.chart import LineChart,Reference
        wb=Workbook();ws=wb.active;ws.title='Bao cao';ws.append(headers)
        for row in values:ws.append([safe(v) for v in row])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='172554');c.font=Font(color='FFFFFF',bold=True)
        for col,key in enumerate(keys,1):
            ws.column_dimensions[ws.cell(1,col).column_letter].width=20 if key!='note' else 35
            for row in range(2,ws.max_row+1):
                cell=ws.cell(row,col)
                if key in ('cost','received','profit','cash','cumulative','opening','pending_cost'):cell.number_format='#,##0;[Red]-#,##0'
                if key=='number':cell.number_format='@'
        metric=next((k for k in ('cumulative','cash','profit','score') if k in keys),None)
        if rows and metric:
            chart=LineChart();chart.title=HEADERS.get(metric,metric);chart.y_axis.title='Điểm' if metric=='score' else 'Đồng'
            col=keys.index(metric)+1;chart.add_data(Reference(ws,min_col=col,min_row=1,max_row=ws.max_row),titles_from_data=True)
            chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row));ws.add_chart(chart,f'A{ws.max_row+4}')
        info=wb.create_sheet('Phuong phap');info.append([title]);info.append([note]);info.column_dimensions['A'].width=110;info['A2'].alignment=Alignment(wrap_text=True)
        buf=io.BytesIO();wb.save(buf);return buf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    if fmt=='pdf':
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.graphics.shapes import Drawing,PolyLine,String,Line
        if 'Viet' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('Viet',str(assets()/'DejaVuSans.ttf')))
        style=ParagraphStyle('body',fontName='Viet',fontSize=8,leading=12)
        title_style=ParagraphStyle('title',parent=style,fontSize=17,leading=23,textColor=colors.HexColor('#172554'))
        buf=io.BytesIO();doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=26,leftMargin=26,topMargin=28,bottomMargin=28)
        story=[Paragraph(escape(title),title_style),Spacer(1,8),Paragraph(escape(note),style),Spacer(1,12)]
        metric=next((k for k in ('cumulative','cash','profit','score') if k in keys),None)
        if len(rows)>1 and metric:
            vals=[float(r.get(metric) or 0) for r in rows];lo=min(vals);hi=max(vals);span=hi-lo or 1
            drawing=Drawing(730,140);pts=[(35+i*680/(len(vals)-1),25+(v-lo)*90/span) for i,v in enumerate(vals)]
            drawing.add(Line(35,25,715,25,strokeColor=colors.grey));drawing.add(PolyLine(pts,strokeColor=colors.HexColor('#2563eb'),strokeWidth=1.5))
            drawing.add(String(35,125,HEADERS.get(metric,metric),fontName='Viet',fontSize=9));drawing.add(String(35,14,str(rows[0].get(keys[0],'')),fontName='Viet',fontSize=7));drawing.add(String(715,14,str(rows[-1].get(keys[0],'')),fontName='Viet',fontSize=7,textAnchor='end'));drawing.add(String(35,3,f'Min {lo:,.0f} • Max {hi:,.0f}',fontName='Viet',fontSize=7));story.append(drawing)
        # Keep rows legible by splitting wide datasets into complementary tables.
        chunks=[list(range(i,min(i+7,len(keys)))) for i in range(0,len(keys),7)]
        for indexes in chunks:
            if indexes[0]!=0:indexes=[0]+indexes
            data=[[Paragraph(escape(headers[i]),style) for i in indexes]]
            for row in values:
                data.append([Paragraph(escape('—' if row[i] is None else f'{row[i]:,}' if isinstance(row[i],(int,float)) else str(row[i])),style) for i in indexes])
            t=Table(data,repeatRows=1,splitInRow=1,colWidths=[(landscape(A4)[0]-64)/len(indexes)]*len(indexes),hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dbeafe')),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fafc')])]))
            story.extend([t,Spacer(1,14)])
        def footer(c,doc):
            c.setFont('Viet',8);c.drawString(26,14,'XSMB V2 PRO • Báo cáo thống kê • Không phải cam kết kết quả');c.drawRightString(815,14,str(doc.page))
        doc.build(story,onFirstPage=footer,onLaterPages=footer)
        return buf.getvalue(),'application/pdf'
    raise ValueError('Chỉ hỗ trợ CSV, Excel và PDF.')
