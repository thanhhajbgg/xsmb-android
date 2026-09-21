import io,json,threading,urllib.request,urllib.error
import pytest
from openpyxl import load_workbook
from src.pro_service import ProService
from src.storage import Repository
from src.pro_server import make_server
from src.reports import export_bytes

@pytest.fixture
def server(tmp_path):
 server=make_server(ProService(Repository(tmp_path/'db')))
 thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
 yield server
 server.shutdown();server.server_close();thread.join()

def request(server,path,data=None,token=True):
 headers={'X-App-Token':server.token} if token else {}
 raw=None if data is None else json.dumps(data).encode()
 return urllib.request.urlopen(urllib.request.Request(f'http://127.0.0.1:{server.server_port}{path}',data=raw,headers=headers))

def test_server_security_and_ledger(server):
 assert b'XSMB' in request(server,'/').read()
 with pytest.raises(urllib.error.HTTPError) as e:request(server,'/api/state',token=False)
 assert e.value.code==403
 row=json.load(request(server,'/api/ledger/save',dict(day='15/09/2026',number='88',points='100',cost_rate='23000',payout_rate='80000',hits='0')))
 assert row['id']>0
 finance=json.load(request(server,'/api/finance'))
 assert finance['profit']==-2300000
 with pytest.raises(urllib.error.HTTPError):request(server,'/api/ledger/save',dict(day='bad'))
 for fmt in ['csv','xlsx','pdf']:
  response=request(server,'/api/export',dict(kind='ledger',format=fmt));assert len(response.read())>50

def test_excel_exact_values_and_formula_safety():
 rows=[dict(day='2026-09-15',number='08',profit=-2300000,note='=1+2')]
 data,_=export_bytes('xlsx','Sổ thắng/thua',rows,['day','number','profit','note'],'Đơn vị: đồng')
 wb=load_workbook(io.BytesIO(data));ws=wb.active
 assert ws['B2'].value=='08' and ws['C2'].value==-2300000
 assert ws['D2'].data_type!='f'

def test_pdf_unicode_and_pages():
 import pymupdf
 data,_=export_bytes('pdf','Sổ thắng/thua',[dict(day='2026-09-15',profit=-2300000)],['day','profit'],'Đơn vị: đồng')
 doc=pymupdf.open(stream=data,filetype='pdf');text=''.join(p.get_text() for p in doc)
 assert 'Sổ thắng/thua' in text and '-2,300,000' in text

def test_server_rejects_foreign_origin(server):
 req=urllib.request.Request(f'http://127.0.0.1:{server.server_port}/api/settings',data=b'{"initial_capital": 9}',headers={'X-App-Token':server.token,'Origin':'https://example.org'})
 with pytest.raises(urllib.error.HTTPError) as e:urllib.request.urlopen(req)
 assert e.value.code==403
 assert server.app.service.get('initial_capital',0)==0

def test_pdf_large_note_and_multiple_pages():
 import pymupdf
 rows=[dict(day='2026-09-15',note='Ghi chú tiếng Việt dài. '*200,profit=5700000) for _ in range(5)]
 data,_=export_bytes('pdf','Nhật ký',rows,['day','profit','note'],'Đơn vị đồng')
 doc=pymupdf.open(stream=data,filetype='pdf');assert len(doc)>1

def test_cau_routes_job_and_exports(server):
 from test_cau import history
 import time
 for d in history(66):server.app.service.repo.save_draw(d)
 result=json.load(request(server,'/api/cau?window=60&min_support=10'))
 assert len(result['pairs'])==20 and result['snapshot']['pair']==result['selected']
 request(server,'/api/job',dict(kind='cau',window=60,min_support=10,points=100)).close()
 deadline=time.monotonic()+10
 while server.app.status()['running'] and time.monotonic()<deadline:time.sleep(.02)
 assert not server.app.status()['running'] and server.app.status()['error'] is None
 assert server.app.status()['result']['days']==6
 for kind in ['cau','cau_backtest']:
  for fmt in ['csv','xlsx','pdf']:
   r=request(server,'/api/export',dict(kind=kind,format=fmt,window=60,min_support=10));assert len(r.read())>50
 assert server.app.service.ledger.rows()==[]
